from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from .models import CONTACT_HEADERS, JOB_HEADERS, Contact, now_iso


def read_contacts(path: str | Path) -> list[Contact]:
    path = Path(path)
    if path.suffix.lower() == ".json":
        rows = json.loads(path.read_text(encoding="utf-8"))
    else:
        with path.open(newline="", encoding="utf-8-sig") as handle:
            rows = list(csv.DictReader(handle))
    return [Contact.from_mapping(row) for row in rows]


def deduplicate_contacts(contacts: Iterable[Contact]) -> list[Contact]:
    result: list[Contact] = []
    seen: set[tuple[str, str]] = set()
    for contact in contacts:
        key = (contact.email.casefold().strip(), contact.company.casefold().strip())
        if key in seen and key != ("", ""):
            continue
        seen.add(key)
        result.append(contact)
    return result


def _setup_sheet(sheet, headers: list[str], title: str) -> None:
    sheet.title = title
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{chr(64 + min(len(headers), 26))}1"
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[chr(64 + index)].width = max(14, min(30, len(header) + 4))


def export_workbook(contacts: Iterable[Contact], output: str | Path, jobs: list[dict] | None = None) -> Path:
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    contacts_sheet = workbook.active
    _setup_sheet(contacts_sheet, CONTACT_HEADERS, "Recruiter Contacts")
    for contact in deduplicate_contacts(contacts):
        row = {
            "Include": contact.include, "Contact Name": contact.contact_name,
            "Job Title": contact.job_title, "Company": contact.company,
            "Email": contact.email, "Phone": contact.phone, "Province": contact.province,
            "City": contact.city, "Recruiter Type": contact.recruiter_type,
            "Target Role Tags": contact.target_role_tags, "Relevant Job Title": contact.relevant_job_title,
            "Job URL": contact.job_url, "Source URL": contact.source_url,
            "Source Type": contact.source_type, "Collected At": contact.collected_at or now_iso(),
            "Email Confidence": contact.email_confidence, "Permission Status": contact.permission_status,
            "Status": contact.status, "Notes": contact.notes, "Last Contacted At": contact.last_contacted_at,
            "Opt Out": contact.opt_out, "Do Not Contact Reason": contact.do_not_contact_reason,
        }
        contacts_sheet.append([row.get(header, "") for header in CONTACT_HEADERS])
    if contacts_sheet.max_row >= 2:
        table = Table(displayName="RecruiterContacts", ref=f"A1:V{contacts_sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        contacts_sheet.add_table(table)
    validation = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    contacts_sheet.add_data_validation(validation)
    validation.add(f"A2:A{max(2, contacts_sheet.max_row)}")

    if jobs:
        job_sheet = workbook.create_sheet()
        _setup_sheet(job_sheet, JOB_HEADERS, "Job Leads")
        for job in jobs:
            job_sheet.append([job.get(header, "") for header in JOB_HEADERS])

    readme = workbook.create_sheet("README")
    readme.append(["Canada Recruiter Outreach Workbook"])
    readme.append(["Review rows manually, set Include to Y, then re-import the workbook."])
    readme.append(["This export does not send email. Unknown or guessed email addresses must not be used."])
    readme.append(["Source URL and Source Type are required for traceability."])
    readme.column_dimensions["A"].width = 110
    workbook.save(output)
    return output


def import_approved(path: str | Path) -> tuple[list[Contact], list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Recruiter Contacts" not in workbook.sheetnames:
        raise ValueError("Workbook is missing the Recruiter Contacts sheet")
    sheet = workbook["Recruiter Contacts"]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    index = {header: position for position, header in enumerate(headers)}
    errors: list[str] = []
    contacts: list[Contact] = []
    for row_number, values in enumerate(rows, start=2):
        include = str(values[index.get("Include", -1)] or "").strip().upper() if index.get("Include") is not None else ""
        if include != "Y":
            continue
        mapped = {header: values[position] for header, position in index.items() if position < len(values)}
        contact = Contact.from_mapping(mapped)
        if not contact.email:
            errors.append(f"Row {row_number}: email is required for approval")
            continue
        if "@" not in contact.email or contact.email.startswith("@") or contact.email.endswith("@"):
            errors.append(f"Row {row_number}: invalid email")
            continue
        if not contact.source_url:
            errors.append(f"Row {row_number}: source URL is required")
            continue
        contacts.append(contact)
    return deduplicate_contacts(contacts), errors
