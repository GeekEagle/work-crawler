from __future__ import annotations

import io
from pathlib import Path

from openpyxl import Workbook, load_workbook

JOB_FIELDS = [
    "id", "company", "title", "location", "salary", "recruiter_email",
    "source", "url", "description", "posted", "status", "notes", "first_seen", "last_seen",
]


def safe_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def export_jobs_xlsx(items: list[dict], review: bool = False) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Job Leads"
    sheet.append(JOB_FIELDS)
    for item in items:
        sheet.append([safe_cell(item.get(field, "")) for field in JOB_FIELDS])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 45)
        sheet.column_dimensions[column[0].column_letter].width = width
    if review:
        manifest = workbook.create_sheet("export_ids")
        for item in items:
            manifest.append([item["id"]])
        manifest.sheet_state = "hidden"
    readme = workbook.create_sheet("README")
    readme.append(["加拿大职位同步结果"])
    readme.append(["数据来自已配置的公开/授权职位来源；请人工核对职位链接和联系方式。"])
    readme.append(["如为待审核 Excel，可删除不需要的职位整行；不要修改 id 或隐藏的 export_ids 表。"])
    readme.append(["导出不会发送邮件。职位按首次发现时间倒序。"])
    readme.column_dimensions["A"].width = 100
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_xlsx(items: list[dict]) -> bytes:
    return export_jobs_xlsx(items, review=True)


def ids_from_reviewed_xlsx(upload) -> tuple[set[str], set[str]]:
    data = upload.getvalue() if hasattr(upload, "getvalue") else upload
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    if "Job Leads" not in workbook.sheetnames or "export_ids" not in workbook.sheetnames:
        raise ValueError("请上传本程序导出的待审核 Excel，并保留 Job Leads 与隐藏的 export_ids 工作表")
    rows = workbook["Job Leads"].iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows, ())]
    if "id" not in headers:
        raise ValueError("Excel 必须保留 id 列")
    index = headers.index("id")
    kept = {str(row[index]).strip() for row in rows if len(row) > index and row[index]}
    exported = {str(row[0]).strip() for row in workbook["export_ids"].iter_rows(values_only=True) if row[0]}
    if not kept <= exported:
        raise ValueError("Excel 中存在不属于原始导出文件的职位 ID")
    return exported, kept


def write_jobs_xlsx(items: list[dict], output: str | Path) -> Path:
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_bytes(export_jobs_xlsx(items))
    return output
