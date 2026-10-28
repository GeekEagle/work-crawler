from openpyxl import load_workbook

from recruiter_outreach.job_excel import write_jobs_xlsx
from recruiter_outreach.job_sources import canadian_it_jobs, parse_boards
from recruiter_outreach.job_storage import all_jobs, connect
from recruiter_outreach.job_sync import SyncConfig, sync_once


def test_sync_detects_new_and_deduplicates(monkeypatch, tmp_path):
    from recruiter_outreach import job_sync

    sample = {"title": "Machine Learning Engineer", "company": "Example", "location": "Toronto, ON, Canada",
              "url": "https://example.com/jobs/1", "source": "Ashby", "posted": "2026-10-27T10:00:00Z"}
    monkeypatch.setattr(job_sync, "fetch_board", lambda kind, company, slug: [sample])
    config = SyncConfig(sources=("boards",), terms=(), cities=(), boards="ashby|Example|example")
    database = tmp_path / "jobs.sqlite3"
    excel = tmp_path / "jobs.xlsx"
    first = sync_once(config, database, excel)
    second = sync_once(config, database, excel)
    assert first["inserted"] == 1
    assert second["inserted"] == 0
    assert second["skipped"] == 1
    with connect(database) as connection:
        records = all_jobs(connection)
    assert len(records) == 1
    assert records[0]["first_seen"] and records[0]["last_seen"]
    sheet = load_workbook(excel, read_only=True)["Job Leads"]
    assert sheet["A2"].value == records[0]["id"]
    assert sheet["H2"].value == sample["url"]


def test_job_excel_escapes_spreadsheet_formulas(tmp_path):
    path = write_jobs_xlsx([{"id": "safe", "company": "=1+1", "title": "AI Engineer"}], tmp_path / "jobs.xlsx")
    sheet = load_workbook(path, read_only=True)["Job Leads"]
    assert sheet["B2"].value == "'=1+1"
    assert sheet["B2"].data_type == "s"


def test_canadian_filter_and_board_validation():
    jobs = [{"title": "Software Engineer", "location": "Canada"},
            {"title": "Software Engineer", "location": "Remote"},
            {"title": "Sales Manager", "location": "Toronto, ON"}]
    assert canadian_it_jobs(jobs) == [jobs[0]]
    assert parse_boards("ashby|Example|example") == [("ashby", "Example", "example")]
from openpyxl import load_workbook

from recruiter_outreach.job_excel import write_jobs_xlsx
from recruiter_outreach.job_storage import all_jobs, connect, fingerprint, upsert_jobs


def test_company_and_title_are_duplicate_key(tmp_path):
    database = tmp_path / "jobs.sqlite3"
    first = {"company": "Example Corp", "title": "AI Engineer", "location": "Toronto, ON", "source": "A"}
    second = {"company": "Example Corp", "title": "AI Engineer", "location": "Vancouver, BC", "source": "B"}
    assert fingerprint(first) == fingerprint(second)
    with connect(database) as connection:
        assert upsert_jobs(connection, [first]) == (1, 0)
        assert upsert_jobs(connection, [second]) == (0, 1)
        rows = all_jobs(connection)
    assert len(rows) == 1
    assert rows[0]["fetched_at"]
    assert rows[0]["source"] == "A"


def test_excel_contains_acquisition_time(tmp_path):
    path = write_jobs_xlsx([{"id": "1", "company": "Example", "title": "AI Engineer", "fetched_at": "2026-10-28T01:02:03+00:00"}], tmp_path / "jobs.xlsx")
    sheet = load_workbook(path, read_only=True)["Job Leads"]
    headers = [cell.value for cell in sheet[1]]
    assert "获取时间" in headers
    assert sheet.cell(2, headers.index("获取时间") + 1).value == "2026-10-28T01:02:03+00:00"


